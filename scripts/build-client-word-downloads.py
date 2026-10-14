from pathlib import Path
import json,re,io,zipfile,hashlib
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parents[1]
out=root/'docs/client-instructions/2026-10-14.2';out.mkdir(parents=True,exist_ok=True)
logo=out/'brand-logo.png'
if not logo.exists():
 import shutil
 shutil.copyfile('/tmp/itinpro-approved-logo.png',logo)
bank=json.loads((root/'lib/portal/instruction-bank.json').read_text())
signature='Ask the authorised issuer to sign supporting letters by hand on its official letterhead, preferably in BLUE INK. From ITIN PRO’s case-handling experience, black-ink handwritten signatures have sometimes been queried as reproduced or digital signatures. Blue ink is our preparation practice, not a universal IRS ink-colour rule. Upload a clear scan and retain the original. Do not type, paste or digitally insert a signature yourself. If a pension provider or brokerage supplies only an electronically signed or system-generated record, request a handwritten letter from its relationship manager first. Upload its response for individual CAA review; acceptance of an alternative is not assumed or guaranteed.'
review='Our CAA and team will review all uploaded documents and request any further documents or clarification. Uploading, payment and draft approval do not submit your application. We will not submit anything without the required signed application, reviewed evidence and original-document checks.'
versions='These instructions and samples are guidance and may need revision. We keep the version issued to your case. If IRS requirements change or the issuer’s wording is unsuitable, we will request an updated version and explain what to replace. Earlier versions remain in your case history.'
fee='The checkout fee covers Form W-7 and CAA work only. Form 1040-NR preparation is separately quoted after document review, with separate engagement and payment terms. No return preparation fee is included in the ITIN payment. Withholding-certificate preparation or specialist work, if needed, is also separately agreed.'

def new_doc(title,branding=True):
 d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
 for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
  st=d.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
  for border in st.element.xpath('.//w:pBdr'):border.getparent().remove(border)
 d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.space_after=Pt(4)
 d.styles['Title'].font.size=Pt(20);d.styles['Title'].paragraph_format.space_after=Pt(9)
 for name in ['Heading 1','Heading 2']:
  d.styles[name].font.size=Pt(13);d.styles[name].paragraph_format.space_before=Pt(10);d.styles[name].paragraph_format.space_after=Pt(5)
 if branding:d.add_paragraph().add_run().add_picture(str(logo),width=Inches(2.3))
 d.add_paragraph(title,'Title');
 if branding:
  p=sec.footer.paragraphs[0];p.add_run('ITIN PRO LTD | www.itinpro.uk | Version 2026-10-14.2 | Page ')
  field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
  for r in p.runs:r.font.size=Pt(8)
 return d

def p(d,text,bold_prefix=None):
 para=d.add_paragraph()
 if bold_prefix and text.startswith(bold_prefix):para.add_run(bold_prefix).bold=True;para.add_run(text[len(bold_prefix):])
 else:para.add_run(text)
 return para

standard=[('Mandatory','Complete your client details in the portal, including the complete foreign address, passport-matching name, date and place of birth, citizenship, prior ITIN and contact details.'),('Mandatory','Upload a clear colour scan of the complete open passport spread, including the photograph/data page and facing top page, with all edges visible. The passport must be valid at submission.'),('Separate original check','We arrange CAA inspection of the original passport or an acceptable issuing-agency-certified copy. An upload is for preparation and does not complete original-document verification. In-person verification is preferred; obtain tracked-post instructions before sending originals.'),('Office record','Provide the proof of address requested in the portal. This is an ITIN PRO office request and is separate from dependent US residency evidence.'),('If applicable','Upload US visa pages and relevant I-20 or DS-2019, US entry/I-94 evidence, and prior IRS rejection or correspondence. Tell us when these do not apply.')]
headings=['Important notes','No letterhead? How to create a professional one in 10 minutes','Letter checklist — what the letter must contain']
source=root.parent/'upload/ITIN_PRO_Letter_Bank_Complete_Package.zip'
with zipfile.ZipFile(source) as source_zip:
 for g in bank:
  d=new_doc('ITIN application document instructions');p(d,f"Category {g['code']}  {g['title']}");p(d,'Read this guide before uploading. Your case-specific portal checklist identifies the applicable evidence; your CAA confirms the application basis.')
  d.add_heading('Step 1 Your category',1);p(d,g['who']);p(d,'Application basis: '+g['basis'])
  d.add_heading('Step 2 Standard documents',1)
  for label,text in standard:p(d,label+'  '+text,label)
  d.add_heading('Step 3 Supporting documents',1)
  if g['code']=='D3':
   p(d,'In the portal, choose ITIN only or ITIN plus a separately quoted refund return. Confirm whether you are the seller or purchaser, whether completion has occurred, the transaction year and the reporting or withholding-certificate purpose. The saved answers determine mandatory uploads.')
   branches=[('Completed sale ITIN only','Provide the signed sale contract or closing statement, completed Form 8288 and applicant-specific Form 8288-A filed by the buyer or escrow agent. Upload remittance proof if available or requested.'),('Withholding certificate','Provide the signed sale or purchase agreement and completed Form 8288-B requesting reduced or eliminated withholding. The adviser confirms the required signatory. Form 8288-B is not a general purchase exemption form.'),('Purchaser reporting','Confirm the seller is foreign for US tax purposes and the actual FIRPTA reporting obligation. Provide the purchase contract, Form 8288 and relevant Form 8288-A. Buying property alone does not establish an ITIN purpose; uncertain cases require adviser review before checkout.'),('Refund return for a completed prior year sale','Provide closing/sale records, Forms 8288 and 8288-A, withholding remittance proof, acquisition costs, improvement records, selling costs and the final signed Form 1040-NR with schedules. If you request preparation, we quote separately after reviewing the uploads.'),('Specific non recognition transaction','Provide the applicable signed notice and transaction evidence. A generic exemption claim is insufficient and your adviser must confirm the precise route.')]
   for title,body in branches:d.add_heading(title,2);p(d,body)
   d.add_heading('When the refund return can be filed',2);p(d,'For a 2026 calendar-year sale, the normal Form 1040-NR refund return is prepared and filed in 2027, after year-end and once the correct final 2026 forms are available. An eligible ITIN exception application can be reviewed earlier. A withholding-certificate or early-refund request is a different procedure and needs separate adviser review; it is not an early annual Form 1040-NR.');p(d,fee)
  else:
   for i in g['items']:p(d,i['level']+'  '+i['instructions']+((' '+i['context']) if i['context'] else ''),i['level'])
   guidance=g['guidance'].split('No letterhead? How to create a professional one in 10 minutes')[0]
   if g['sample']:d.add_heading('Letter checklist',2)
   for line in guidance.splitlines():
    if not line.strip():continue
    if line in headings:d.add_heading(line.replace('—','').replace('?',''),2)
    else:p(d,line)
  if g['sample']:p(d,'Use the editable Word letter template. Add the issuer’s logo, registered legal name, official address and relevant contact details, or use its existing official letterhead. Replace every highlighted example with accurate information. Send the draft through the portal for checking before it is printed and signed.')
  d.add_heading('Signatures and originals',1);p(d,signature)
  d.add_heading('Review and updated versions',1);p(d,review);p(d,versions)
  p(d,'Upload via your private case at www.itinpro.uk/portal. For upload help, contact Shan@itinpro.uk to request secure access. Do not email identity documents. We will arrange original checks and tell you how to sign the final W-7.')
  if g['code']=='D3':p(d,'Official sources: IRS Instructions for Form W-7; IRS ITIN guidance for foreign property buyers and sellers; IRS Instructions for Form 1040-NR; IRS IRM 21.8.5.4.11. Your adviser checks current requirements before submission.')
  review_start=next(i for i,p in enumerate(d.paragraphs) if p.text=='Review and updated versions')
  for paragraph in d.paragraphs[review_start:-1]:paragraph.paragraph_format.keep_with_next=True
  d.paragraphs[-1].paragraph_format.keep_together=True
  d.save(out/(g['code']+'-instructions.docx'))
  if not g['sample']:continue
  text=g['sample'].split('\n',1)[1]
  filename=next(n for n in source_zip.namelist() if '/'+g['code']+'_' in n)
  with zipfile.ZipFile(io.BytesIO(source_zip.read(filename))) as inner:
   xml=etree.fromstring(inner.read('word/document.xml'));ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
   highlighted=list(dict.fromkeys(xml.xpath('//w:r[w:rPr/w:highlight]//w:t/text()',namespaces=ns)))
  for h in sorted(highlighted,key=len,reverse=True):
   if not h.strip() or h not in text:continue
   if h=='2026':replacement='[TAX YEAR]'
   elif h=='14 October 2026':replacement='[SIGNATURE DATE]'
   elif 'DANIEL ROBERT EXAMPLE' in h:replacement='[APPLICANT FULL LEGAL NAME]'
   elif re.fullmatch(r'\d{2}-\d{7}',h.strip()):replacement='[ISSUER EIN]'
   elif re.search(r'\(\d{3}\)|555-',h):replacement='[US PHONE IF HELD]'
   elif any(x in h.lower() for x in ['road','street','avenue','close','boulevard','lane']):replacement='[RELEVANT ADDRESS]'
   elif 'Sample' in h:replacement='[AUTHORISED SIGNER NAME]'
   elif any(word in h for word in ['Manager','Director','Specialist','Officer','Executor','Member']):replacement='[AUTHORISED SIGNER TITLE]'
   elif 'EIN' in h:replacement='[ISSUER EIN]'
   elif '@' in h:replacement='[ISSUER EMAIL]'
   elif 'Reference' in h:replacement='[ISSUER REFERENCE]'
   elif any(x in h.lower() for x in ['road','street','avenue','close','boulevard','lane']):replacement='[RELEVANT ADDRESS]'
   elif 'Plan Sponsor' in h:replacement='[PLAN SPONSOR NAME]'
   elif 'Plan' in h or '401(k)' in h:replacement='[RETIREMENT PLAN NAME]'
   elif len(h)<5:replacement='[ISSUER NAME]'
   elif any(x in h for x in ['Inc.','Ltd','LLC','Bank','Services','Estate','Foundation','Partners','University','Casino']):replacement='[ISSUER LEGAL NAME]'
   else:replacement='[REPLACE WITH ACTUAL DETAILS]'
   text=text.replace(h,replacement)
  text=text.replace('14 October 2026','[SIGNATURE DATE]').replace('2026','[TAX YEAR]')
  text=re.sub(r'\b\d{2}-\d{7}\b','[ISSUER EIN]',text);text=re.sub(r'\(\d{3}\) 555-\d{4}','[US PHONE IF HELD]',text);text=re.sub(r'[\w.+-]+@[\w.-]*example\b','[ISSUER EMAIL]',text)
  text=re.sub(r'US\$[\d,]+','[ACTUAL AMOUNT]',text)
  text=re.sub(r'\[Instruction — do not copy\][^\n]*','',text)
  text=re.sub(r'(?m)^([A-Za-z ()/\-]{2,50}:)\n',r'\1 ',text)
  text=text.replace('Internal Revenue Service\nITIN Operation','Internal Revenue Service  ITIN Operation').replace('Austin, TX 78714-9342\nUnited States of America','Austin, TX 78714-9342  United States of America')
  if g['code']=='D4':
   text='[SIGNATURE DATE]\nInternal Revenue Service  ITIN Operation\nP.O. Box 149342\nAustin, TX 78714-9342  United States of America\nRe: Form W-7 Exception 5 Treasury Decision 9363\nTo Whom It May Concern:\n[FOREIGN CORPORATION LEGAL NAME], incorporated in [COUNTRY OF INCORPORATION], with its registered office at [ACTUAL REGISTERED OFFICE ADDRESS] and EIN [CORPORATION EIN OR EXPLAIN APPLICATION STATUS], confirms that [APPLICANT FULL LEGAL NAME], who serves as [APPLICANT POSITION], is designated as the person responsible for the corporation’s IRS information-reporting compliance under Treasury Decision 9363.\nThe corporation is required to electronically file [APPLICABLE IRS RETURNS AND REPORTING OBLIGATION] for [TAX YEAR]. An ITIN is needed for the applicant to carry out the specific reporting-compliance responsibility described above.\n[DESCRIBE THE SIGNER’S AUTHORITY TO CONFIRM THIS DESIGNATION. IF THE APPLICANT IS THE SOLE DIRECTOR AND SIGNS, STATE SOLE-DIRECTOR AUTHORITY.]\nWe request that this letter be considered with the applicant’s Form W-7 under Exception 5 T.D. 9363. If further information is required, please contact [AUTHORISED CONTACT NAME] at [CORPORATION EMAIL].\nYours faithfully,\nFor and on behalf of [FOREIGN CORPORATION LEGAL NAME]\n_________________________________\nOriginal handwritten signature in BLUE INK\n[AUTHORISED SIGNER NAME]\n[AUTHORISED SIGNER TITLE]\nDate: [SIGNATURE DATE]'
  text=re.sub(r'\bRobert James Example\b','[DECEASED FULL LEGAL NAME]',text)
  if g['code']=='B14':
   text=text.replace('We act for [AUTHORISED SIGNER TITLE]','We act for [EXECUTOR FULL LEGAL NAME]').replace('from the [ISSUER LEGAL NAME]','from the [ESTATE LEGAL NAME]').replace('For and on behalf of [AUTHORISED SIGNER TITLE]','For and on behalf of [ESTATE LEGAL NAME]')
  # Native editable letter with issuer letterhead placeholders and highlighted fields.
  d=new_doc('Editable supporting letter',False)
  d.sections[0].top_margin=Inches(.5);d.sections[0].bottom_margin=Inches(.5);d.styles['Title'].font.size=Pt(17)
  para=p(d,'Replace highlighted fields with accurate details. Add the issuer’s logo and official letterhead, or use its own letterhead. Delete this heading and instruction paragraph before printing. Ask your CAA to check the draft before obtaining the required handwritten signature, preferably in blue ink.')
  for r in para.runs:r.font.size=Pt(9)
  for line in ['[ISSUER LOGO IF USED]','[ISSUER LEGAL NAME]','[ISSUER OFFICIAL ADDRESS]','[ISSUER EMAIL AND US PHONE IF HELD]','']+text.splitlines():
   if not line.strip():continue
   para=d.add_paragraph();para.paragraph_format.space_after=Pt(0)
   for fragment in re.split(r'(\[[^\]]+\])',line):
    r=para.add_run(fragment)
    if fragment.startswith('['):r.font.highlight_color=WD_COLOR_INDEX.YELLOW
  d.save(out/(g['code']+'-letter.docx'))
hashes={str(p.relative_to(root/'docs/client-instructions')):hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(out.glob('*.docx'))}
(root/'lib/portal/word-download-hashes.json').write_text(json.dumps(hashes,indent=2)+'\n')
print('Created',len(hashes),'editable Word downloads')
